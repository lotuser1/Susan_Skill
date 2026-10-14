#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把知识库里的本地课件（pptx / docx / pdf）转成 markdown，并入知识库可被检索。

设计原则：
  - **原件一律不动**，只生成同名 .md 放在源文件旁边
  - 输出带 front-matter，风格与听悟转写稿一致，便于 dbs-knowledge 统一索引
  - 文件名加「【课件】」前缀，Finder 里一眼能认出
  - 幂等：已转过的默认跳过（--force 可重转）

用法：
    python3 scripts/convert_docs.py --dry        # 演练：列出将转哪些、多少字
    python3 scripts/convert_docs.py --apply      # 真正转换
"""
import os
import argparse
import re
import sys
import time
from pathlib import Path

KB_ROOT = Path(os.environ.get("KB_ROOT", "./知识库"))
SKIP_DIR_PREFIX = "_"
EXTS = {".pptx", ".docx", ".pdf"}
PREFIX = "【课件】"


def clean(s):
    """压掉多余空白，保留换行结构。"""
    s = re.sub(r"[ \t\u3000]+", " ", s or "")
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def from_pptx(fp):
    from pptx import Presentation
    prs = Presentation(str(fp))
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        chunks = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = clean(sh.text_frame.text)
                if t:
                    chunks.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [" | ".join(c.text.strip() for c in r.cells)
                        for r in sh.table.rows]
                if rows:
                    chunks.append("\n".join(rows))
        # 演讲者备注（常藏着真正的讲稿）
        try:
            if slide.has_notes_slide:
                nt = clean(slide.notes_slide.notes_text_frame.text)
                if nt:
                    chunks.append("【备注】" + nt)
        except Exception:  # noqa: BLE001
            pass
        if chunks:
            pages.append((i, "\n\n".join(chunks)))
    return pages, len(prs.slides)


def from_docx(fp):
    import docx
    d = docx.Document(str(fp))
    chunks = []
    for p in d.paragraphs:
        t = clean(p.text)
        if t:
            chunks.append(t)
    for tb in d.tables:
        for r in tb.rows:
            line = " | ".join(c.text.strip() for c in r.cells)
            if line.strip(" |"):
                chunks.append(line)
    body = "\n\n".join(chunks)
    return ([(1, body)] if body else []), 1


def from_pdf(fp):
    import pdfplumber
    pages = []
    with pdfplumber.open(str(fp)) as pdf:
        for i, pg in enumerate(pdf.pages, 1):
            t = clean(pg.extract_text() or "")
            if t:
                pages.append((i, t))
        total = len(pdf.pages)
    return pages, total


def render(fp, pages, total_pages, src_type):
    n_chars = sum(len(t) for _, t in pages)
    title = fp.stem
    fm = [
        "---",
        f'title: "{title}"',
        f'source: "本地课件"',
        f'source_file: "{fp.name}"',
        f'source_type: "{src_type}"',
        f'pages: {total_pages}',
        f'word_count: {n_chars}',
        f'converted_at: "{time.strftime("%Y-%m-%d %H:%M:%S")}"',
        "---", "",
        f"# {title}", "",
        f"> 来源：本地课件（{src_type.upper()}）· {total_pages} 页 · 约 {n_chars:,} 字", "",
    ]
    body = []
    for i, t in pages:
        body.append(f"## 第 {i} 页\n")
        body.append(t)
        body.append("")
    return "\n".join(fm + body)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--force", action="store_true", help="已转过的也重转")
    a = ap.parse_args()
    if not (a.dry or a.apply):
        ap.print_help()
        return 1

    files = []
    for fp in sorted(KB_ROOT.rglob("*")):
        if not fp.is_file() or fp.suffix.lower() not in EXTS:
            continue
        rel = fp.relative_to(KB_ROOT)
        if any(p.startswith(SKIP_DIR_PREFIX) for p in rel.parts):
            continue
        if fp.name.startswith("~$"):  # Office 临时文件
            continue
        files.append(fp)

    print(f"待转换 {len(files)} 个文件（原件保持不动）\n")
    t0 = time.time()
    done = skipped = failed = 0
    total_chars = 0
    for fp in files:
        out = fp.parent / f"{PREFIX}{fp.stem}.md"
        ext = fp.suffix.lower()
        try:
            if out.exists() and not a.force:
                print(f"  — 已存在，跳过：{out.name}")
                skipped += 1
                continue
            if ext == ".pptx":
                pages, tp = from_pptx(fp)
            elif ext == ".docx":
                pages, tp = from_docx(fp)
            else:
                pages, tp = from_pdf(fp)
            n = sum(len(t) for _, t in pages)
            if not pages:
                print(f"  ⚠️ 无文字内容，跳过：{fp.name}")
                skipped += 1
                continue
            md = render(fp, pages, tp, ext.lstrip("."))
            print(f"  {'[演练] ' if a.dry else ''}{fp.name}")
            print(f"          → {out.name} · {tp} 页 · {n:,} 字")
            if a.apply:
                out.write_text(md, encoding="utf-8")
            total_chars += n
            done += 1
        except Exception as e:  # noqa: BLE001
            print(f"  ❌ 失败：{fp.name} — {type(e).__name__}: {e}")
            failed += 1

    print(f"\n合计：转换 {done} · 跳过 {skipped} · 失败 {failed} · "
          f"提取 {total_chars:,} 字 · 耗时 {time.time() - t0:.0f}s")
    if a.dry:
        print("（演练模式，未写任何文件。加 --apply 执行）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
